from fastapi import FastAPI, HTTPException, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import List, Dict, Optional
from datetime import datetime
from docx import Document
import io
import zipfile
import PyPDF2
from fastapi.responses import StreamingResponse

# --- REAL WORLD INTEGRATIONS ---
import json
import os
import random
import requests
from groq import Groq
from supabase import create_client, Client

# 1. Groq Setup
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")
client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None
MODEL_ID = "llama-3.3-70b-versatile"

# 2. Supabase Setup
SUPABASE_URL = os.environ.get("SUPABASE_URL")
SUPABASE_KEY = os.environ.get("SUPABASE_KEY")
try:
    supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY) if SUPABASE_URL and SUPABASE_KEY else None
except Exception as e:
    print(f"Supabase Init Error: {e}")
    supabase = None

# =====================================================================
# SCHEMAS
# =====================================================================

class StructuredCriteria(BaseModel):
    age_min: Optional[int] = None
    age_max: Optional[int] = None
    conditions_required: List[str] = []
    conditions_excluded: List[str] = []

class Protocol(BaseModel):
    protocol_id: str
    nct_id: str
    title: str
    indication: str
    phase: str
    target_enrollment: int
    inclusion_criteria: List[str]
    exclusion_criteria: List[str]
    structured_criteria: StructuredCriteria
    geographies: List[str]
    created_at: str

class FilterRequest(BaseModel):
    protocol: Protocol

class SiteRankingRequest(BaseModel):
    protocol: Protocol

class SimulationRequest(BaseModel):
    protocol: Protocol
    selected_site_ids: List[str]

class DocGenerationRequest(BaseModel):
    protocol: Protocol
    selected_site_ids: List[str]
    doc_types: List[str]


# =====================================================================
# GLOBAL STATE (Caches dynamic sites during the demo session)
# =====================================================================
db_sites = []

STATE_ABBR_MAP = {
    "California": "CA", "New York": "NY", "Texas": "TX", "Florida": "FL",
    "Virginia": "VA", "Maryland": "MD", "North Carolina": "NC", "District of Columbia": "DC",
    "Massachusetts": "MA"
}
REVERSE_STATE_MAP = {v: k for k, v in STATE_ABBR_MAP.items()}
DEMO_TARGET_STATES = ["CA", "NY", "TX", "FL", "VA", "MD", "NC", "DC"]

# =====================================================================
# FASTAPI ENDPOINTS
# =====================================================================

app = FastAPI(title="Clinical Trial Start-Up API (Live Data Edition)")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

@app.post("/protocol/parse")
async def parse_protocol(file: UploadFile = File(...)):
    try:
        if not client: raise HTTPException(status_code=500, detail="Groq API Key missing.")

        pdf_reader = PyPDF2.PdfReader(file.file)
        extracted_text = "".join([page.extract_text() for page in pdf_reader.pages])[:15000]
            
        system_prompt = """
        You are an expert clinical trial data extraction AI. 
        Determine if the text is genuinely a Clinical Trial Protocol.
        
        Required JSON structure:
        {
            "is_valid_protocol": <boolean>,
            "rejection_reason": "<string>",
            "protocol_id": "Generate a random ID like P-102",
            "nct_id": "Extract NCT ID or use 'UNKNOWN'",
            "title": "Extract full study title",
            "indication": "Extract the primary disease or condition in lowercase",
            "phase": "Extract trial phase",
            "target_enrollment": <integer>,
            "inclusion_criteria": ["criteria 1"],
            "exclusion_criteria": ["criteria 1"],
            "structured_criteria": {
                "age_min": <integer or null>,
                "age_max": <integer or null>,
                "conditions_required": ["disease 1"],
                "conditions_excluded": ["disease 1"]
            },
            "geographies": ["VA", "MD", "DC", "CA", "NY", "TX", "NC", "FL"],
            "created_at": "ISO 8601 Timestamp string"
        }
        """
        
        chat_completion = client.chat.completions.create(
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": f"Protocol Text:\n{extracted_text}"}
            ],
            model=MODEL_ID,
            response_format={"type": "json_object"}
        )
        
        data = json.loads(chat_completion.choices[0].message.content)

        if data.get("is_valid_protocol") is False:
            raise ValueError(data.get("rejection_reason", "Document rejected."))
            
        data["title"] = f"[Groq Extracted] {data.get('title', 'Unknown Title')}"
        if not data.get("geographies"):
             data["geographies"] = DEMO_TARGET_STATES
             
        return data

    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error parsing document: {str(e)}")

@app.post("/participants/filter")
async def filter_participants(req: FilterRequest):
    if not supabase:
        raise HTTPException(status_code=500, detail="Supabase integration missing. Check environment variables.")

    try:
        # FORCE GEOGRAPHIC FENCE: Always use our 8 demo states so the heatmap works beautifully
        geos_full_names = [REVERSE_STATE_MAP.get(g, g) for g in DEMO_TARGET_STATES]
        
        # 1. Live Supabase Query - Try uppercase STATE first, then fallback to lowercase state
        try:
            response = supabase.table('patients').select('*').in_('STATE', geos_full_names).execute()
        except Exception:
            response = supabase.table('patients').select('*').in_('state', geos_full_names).execute()

        raw_patients = response.data

        # 2. Local Age Processing & Re-mapping abbreviations
        site_counts = {state: 0 for state in DEMO_TARGET_STATES}
        eligible_patients = []
        criteria = req.protocol.structured_criteria

        current_year = datetime.now().year

        for pt in raw_patients:
            # Case-insensitive dictionary lookups for Supabase columns
            pt_state = pt.get('state') or pt.get('STATE') or ''
            pt_gender = pt.get('gender') or pt.get('GENDER') or 'U'
            pt_birthdate = pt.get('birthdate') or pt.get('BIRTHDATE') or ''
            pt_id = pt.get('id') or pt.get('Id') or pt.get('ID')

            try:
                birth_year = int(pt_birthdate[:4])
                age = current_year - birth_year
            except:
                age = 45 # Fallback

            if criteria.age_min and age < criteria.age_min: continue
            if criteria.age_max and age > criteria.age_max: continue

            pt_state_abbr = STATE_ABBR_MAP.get(pt_state, pt_state)
            
            # Only count them if they are in our target states
            if pt_state_abbr in site_counts:
                eligible_patients.append({
                    "patient_id": pt_id,
                    "demographics": {"age": age, "gender": pt_gender},
                    "location": {"state": pt_state_abbr}
                })
                site_counts[pt_state_abbr] += 1

        return {
            "total_eligible": len(eligible_patients),
            "distribution_by_state": site_counts,
            "sample": eligible_patients[:5]
        }
    except Exception as e:
        print("Supabase Query Error:", str(e))
        raise HTTPException(status_code=500, detail="Failed to query Supabase database.")

@app.post("/sites/rank")
async def rank_sites(req: SiteRankingRequest):
    global db_sites
    try:
        if not client: raise HTTPException(status_code=500, detail="Groq API Key missing.")
        
        # 1. LIVE API PING: Fetch real hospitals from clinicaltrials.gov
        url = f"https://clinicaltrials.gov/api/v2/studies?query.cond={req.protocol.indication}&pageSize=100"
        ct_response = requests.get(url).json()
        
        real_sites = []
        seen_names = set()
        
        for study in ct_response.get("studies", []):
            locations = study.get("protocolSection", {}).get("contactsLocationsModule", {}).get("locations", [])
            for loc in locations:
                # GEOGRAPHIC FENCE: Strictly filter for United States
                country = loc.get("country", "")
                if country != "United States":
                    continue

                name = loc.get("facility", "")
                if not name or name in seen_names: continue
                
                # Format to our schema
                state_raw = loc.get("state", "")
                state_abbr = STATE_ABBR_MAP.get(state_raw, state_raw) 

                # GEOGRAPHIC FENCE: Only include sites in our 8 target states so they hit the heatmap!
                if state_abbr not in DEMO_TARGET_STATES:
                    continue

                seen_names.add(name)

                real_sites.append({
                    "site_id": f"CT-{len(real_sites)+1000}",
                    "npi": str(random.randint(1000000000, 9999999999)),
                    "name": name,
                    "organization_type": "Research Institution",
                    "specialty": req.protocol.indication.title(),
                    "location": {"city": loc.get("city", "Unknown"), "state": state_abbr, "zip": loc.get("zip", "00000")},
                    "therapeutic_areas": [req.protocol.indication.lower()],
                    "metrics": {
                        "past_enrollment_rate": round(random.uniform(2.0, 15.0), 1),
                        "activation_time_days": random.randint(30, 90),
                        "patient_pool_size": random.randint(1000, 8000),
                        "performance_score": 0.0
                    },
                    "capabilities": {"has_trial_experience": True, "staff_count": random.randint(10, 80)}
                })

        # Fallback just in case CT.gov returns zero US matches for a weird indication
        if not real_sites:
            real_sites.append({
                "site_id": "CT-FALLBACK", "npi": "1234567890", "name": "Mayo Clinic (Demo Fallback)",
                "organization_type": "Research Institution", "specialty": req.protocol.indication.title(),
                "location": {"city": "Jacksonville", "state": "FL", "zip": "32224"},
                "therapeutic_areas": [req.protocol.indication.lower()],
                "metrics": {"past_enrollment_rate": 12.5, "activation_time_days": 45, "patient_pool_size": 5000, "performance_score": 0.0},
                "capabilities": {"has_trial_experience": True, "staff_count": 45}
            })

        db_sites = real_sites 

        # 2. AI RANKING
        compressed = [{"id": s["site_id"], "state": s["location"]["state"], "rate": s["metrics"]["past_enrollment_rate"]} for s in db_sites[:30]]

        system_prompt = """You are an AI Clinical Site Selection Engine.
        Return ONLY valid JSON:
        {"top_sites": [{"site_id": "string", "score": <float 0.0-1.0>, "breakdown": {"enrollment_rate_component": <float>, "patient_availability_component": <float>, "therapeutic_match_component": <float>, "geography_match_component": <float>}}]}
        """

        user_prompt = f"Target Indication: {req.protocol.indication}\nGeographies: {DEMO_TARGET_STATES}\nSites: {json.dumps(compressed)}"

        chat = client.chat.completions.create(
            messages=[{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
            model=MODEL_ID, response_format={"type": "json_object"}
        )
        llm_results = json.loads(chat.choices[0].message.content)

        ranked = []
        for item in llm_results.get("top_sites", []):
            original = next((s for s in db_sites if s["site_id"] == item["site_id"]), None)
            if original:
                copy = original.copy()
                copy["metrics"]["performance_score"] = item["score"]
                ranked.append({"site": copy, "score": item["score"], "breakdown": item.get("breakdown", {})})

        ranked.sort(key=lambda x: x["score"], reverse=True)
        return {"top_sites": ranked}

    except Exception as e:
        print(f"Ranking Error: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/enrollment/simulate")
async def simulate_enrollment(req: SimulationRequest):
    try:
        global db_sites
        selected = [s for s in db_sites if s["site_id"] in req.selected_site_ids]
        if not selected: raise HTTPException(status_code=400, detail="No sites selected")

        target = req.protocol.target_enrollment or 100
        compressed_sites = [{"site_id": s["site_id"], "avg_monthly_enrollment": s["metrics"]["past_enrollment_rate"], "activation_days": s["metrics"]["activation_time_days"]} for s in selected]

        system_prompt = """Return ONLY valid JSON.
        {"estimated_completion_month": <integer>, "timeline": [{"month": <integer>, "monthly_enrolled": <integer>, "cumulative_enrolled": <integer>}]}
        """

        user_prompt = f"Target Total Enrollment: {target}\nSelected Sites Data:\n{json.dumps(compressed_sites)}\nSimulate timeline."

        chat_completion = client.chat.completions.create(
            messages=[{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
            model=MODEL_ID, response_format={"type": "json_object"}
        )
        
        llm_timeline = json.loads(chat_completion.choices[0].message.content)

        return {
            "protocol_id": req.protocol.protocol_id,
            "estimated_completion_month": llm_timeline.get("estimated_completion_month", 0),
            "total_sites_active": len(selected),
            "timeline": llm_timeline.get("timeline", [])
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/documents/generate")
async def generate_document(req: DocGenerationRequest):
    global db_sites
    selected = [s for s in db_sites if s["site_id"] in req.selected_site_ids]
    zip_buffer = io.BytesIO()
    
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for site in selected:
            if "FDA_1572" in req.doc_types:
                doc = Document()
                doc.add_heading('FDA Form 1572', 0)
                doc.add_paragraph(f"NCT ID: {req.protocol.nct_id}")
                doc.add_paragraph(f"Indication: {req.protocol.indication.title()}")
                doc.add_paragraph(f"Site: {site['name']} (NPI: {site['npi']})")
                f = io.BytesIO()
                doc.save(f)
                zip_file.writestr(f"FDA_1572_{site['npi']}.docx", f.getvalue())
                
            if "CDA" in req.doc_types:
                doc = Document()
                doc.add_heading('CDA', 0)
                doc.add_paragraph(f"Agreement between Sponsor and {site['name']}.")
                f = io.BytesIO()
                doc.save(f)
                zip_file.writestr(f"CDA_{site['npi']}.docx", f.getvalue())

    zip_buffer.seek(0)
    return StreamingResponse(zip_buffer, media_type="application/zip", headers={"Content-Disposition": f"attachment; filename=Docs_{req.protocol.protocol_id}.zip"})